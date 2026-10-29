# 金額が、出典の資料の「掲載ページ」に本当に書かれているかを照合するプログラム
#
#   python -I tools/verify_pages.py <資料PDFを置いたフォルダ> [自治体コード …] [--render <画像の保存先>]
#
# ・資料PDFは tools/fetch_sources.py で取ってくる（<フォルダ>/<自治体コード>/<出典ID>_<URLの最後の部分> に置かれる）
#   （例：221309-S01 の https://～/22_setumeisho08.pdf → <フォルダ>/221309/221309-S01_22_setumeisho08.pdf）
# ・エクセル（.xlsx）の資料は、掲載ページに「XLSX「シート名」シート 94行目」のように書くと、その行のセルを照合します
# ・予算・決算（data/budgets）、財源（data/funds）、支払先（data/recipients）の各行について、
#   「掲載ページ」の文字の中に「資料の数字」（と財源の数字）があるかを調べ、結果を表にして出します
# ・紙をスキャンしたPDFなど、文字を取り出せないページは「目で確認が必要」と出ます。
#   --render を付けると、そのページを画像（PNG）にして保存するので、画像を見て確かめてください
# ・このプログラムは「確認」の列を書きかえません。照合できた行を人が確かめて「確認済み」にします
#
# 注意：PDF はインターネットから取ったファイルなので、このプログラムは -I を付けて、PDFのフォルダの外から実行してください。

import csv
import glob
import os
import re
import sys
import unicodedata
from urllib.parse import unquote, urlparse

import pymupdf

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def norm(s):
    s = unicodedata.normalize("NFKC", s or "")
    # 見えない空白（ゼロ幅スペースなど）が文字の間に入っているPDFもあるので取り除く
    return re.sub(r"[\s,，、\u200b\u200c\u200d\u2060\ufeff]", "", s)


def read_csvs(table):
    rows = []
    for fp in sorted(glob.glob(os.path.join(ROOT, "data", table, "*.csv"))):
        with open(fp, encoding="utf-8-sig") as f:
            for i, r in enumerate(csv.DictReader(f), start=2):
                r["_file"], r["_line"] = os.path.relpath(fp, ROOT), i
                rows.append(r)
    return rows


def pages_of(page_field):
    """「PDF12ページ目（印刷10）」「p.12」「12-13」などから、PDFのページ番号（1から）を取り出す"""
    s = unicodedata.normalize("NFKC", page_field or "")
    m = re.search(r"PDF\s*(?:の)?\s*(\d+)(?:\s*[-~〜～]\s*(\d+))?", s)
    if not m:
        m = re.search(r"(\d+)(?:\s*[-~〜～]\s*(\d+))?", s)
    if not m:
        return []
    a = int(m.group(1))
    b = int(m.group(2)) if m.group(2) else a
    return list(range(a, min(b, a + 5) + 1))


def raw_variants(raw):
    """資料の数字の、ページの文字の中での書かれ方の候補"""
    t = norm(raw).replace("円", "")
    t = re.sub(r"^[△▲-]", "", t)
    out = {t}
    m = re.fullmatch(r"(\d+)億(\d+)", t)
    if m:
        out.add(m.group(1) + m.group(2).zfill(4))
    return [v for v in out if v]


def judge(main_text, raws, r):
    """数字（と「資料の記載」の言葉）が main_text にあるか。数字が足りないときは（""、足りない数字）を返す"""
    miss = [x for x in raws if not any(v in main_text for v in raw_variants(x))]
    if miss:
        return "", miss
    # 「資料の記載」（原文の引用）の中の言葉（数字以外の4文字以上）が、同じページにあるか
    words = [w for w in re.split(r"[\d.%()（）「」『』・:：\-－/／|｜]+", norm(r.get("資料の記載", ""))) if len(w) >= 4]
    if not words or any(w in main_text for w in words):
        return "OK", []
    return "数字はあるが、資料の記載の言葉がページに見つからない（目で確認）", []


XLSX = {}


def xlsx_text(fp, page_field):
    """「XLSX「予算事業一覧」シート 94行目（下段の所要一般財源は95行目）」のような掲載ページから、その行のセルの文字をつなげて返す"""
    import openpyxl

    s = unicodedata.normalize("NFKC", page_field or "")
    rows = [int(n) for n in re.findall(r"(\d+)\s*行目", s)]
    if not rows:
        return ""
    if fp not in XLSX:
        XLSX[fp] = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    wb = XLSX[fp]
    m = re.search(r"「(.+?)」\s*シート", s)
    ws = wb[m.group(1)] if m and m.group(1) in wb.sheetnames else wb.worksheets[0]
    cells = []
    for n in rows:
        for row in ws.iter_rows(min_row=n, max_row=n, values_only=True):
            for v in row:
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                if v is not None:
                    cells.append(str(v))
    return norm(" ".join(cells))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    render_dir = None
    if "--render" in args:
        i = args.index("--render")
        render_dir = args[i + 1]
        del args[i:i + 2]
    if not args:
        sys.exit(__doc__ or "使い方：python -I tools/verify_pages.py <資料PDFのフォルダ> [自治体コード …]")
    dl, codes = args[0], set(args[1:])

    sources = {r["出典ID"]: r for r in read_csvs("sources")}
    docs = {}

    def doc_for(sid):
        s = sources.get(sid)
        if not s:
            return None, "出典なし"
        name = os.path.basename(unquote(urlparse(s["出典URL"]).path))
        fp = os.path.join(dl, s["自治体コード"], f"{sid}_{name}")
        if not os.path.exists(fp):
            return None, f"PDFが手元にない（{name}）"
        if fp.lower().endswith(".xlsx"):
            return fp, ""
        # URLが「.pdf」で終わらない資料もあるので、ファイルの先頭でPDFかどうかを見る
        with open(fp, "rb") as f:
            if f.read(5) != b"%PDF-":
                return None, f"PDFが手元にない（{name}）"
        if fp not in docs:
            docs[fp] = pymupdf.open(fp, filetype="pdf")
        return docs[fp], ""

    results = []
    checks = []
    for r in read_csvs("budgets"):
        checks.append(("予算・決算", r, [r["資料の数字"]]))
    for r in read_csvs("funds"):
        nums = [r[k] for k in ["国庫支出金（円）", "都道府県支出金（円）", "国・県支出金（区別なし）（円）", "地方債（円）", "その他財源（円）", "一般財源（円）"] if r[k]]
        checks.append(("財源", r, [str(int(n) // 1000) if int(n) % 1000 == 0 else n for n in nums]))
    for r in read_csvs("recipients"):
        if r["資料の数字"]:
            checks.append(("支払先", r, [r["資料の数字"]]))

    for kind, r, raws in checks:
        if codes and r["自治体コード"] not in codes:
            continue
        doc, why = doc_for(r["出典ID"])
        pg = pages_of(r["掲載ページ"])
        res = ""
        if not doc:
            res = why
        elif isinstance(doc, str):  # エクセルの資料：掲載ページの「シート名」と「○行目」の行を見る
            main_text = xlsx_text(doc, r["掲載ページ"])
            if not main_text:
                res = "掲載ページのシート名・行番号が読めない"
            else:
                res, miss = judge(main_text, raws, r)
                res = res or f"数字が掲載の行にない（{'・'.join(miss)}）"
            doc = None
        elif not pg:
            res = "掲載ページの番号が読めない"
        else:
            texts = {}
            for p in set(pg + [pg[0] - 1, pg[-1] + 1]):
                if 1 <= p <= doc.page_count:
                    # 段組みのページでは数字が途中で切れて取り出されることがあるので、位置順に並べ直した文字も足す
                    texts[p] = norm(doc[p - 1].get_text()) + "|" + norm(doc[p - 1].get_text(sort=True))
            main_text = "".join(texts.get(p, "") for p in pg)
            if len(main_text) < 30:
                res = "目で確認が必要（文字を取り出せないページ）"
            else:
                res, miss = judge(main_text, raws, r)
                if not res:
                    near = [p for p, t in texts.items() if p not in pg and all(any(v in t for v in raw_variants(x)) for x in miss)]
                    res = f"数字が掲載ページにない（{'・'.join(miss)}）" + (f"→ PDF{near[0]}ページにある" if near else "")
            if res != "OK" and render_dir and doc and pg:
                os.makedirs(render_dir, exist_ok=True)
                for p in pg[:2]:
                    if 1 <= p <= doc.page_count:
                        out = os.path.join(render_dir, f"{r['自治体コード']}_{os.path.basename(doc.name)}_p{p}.png")
                        if not os.path.exists(out):
                            doc[p - 1].get_pixmap(dpi=110).save(out)
        results.append([kind, r["_file"], r["_line"], r.get("事業ID", ""), r.get("年度", ""), r.get("段階", ""), "・".join(raws), r["出典ID"], r["掲載ページ"], res])

    out = os.path.join(ROOT, "tools", "verify_report.csv")
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["種類", "ファイル", "行", "事業ID", "年度", "段階", "資料の数字", "出典ID", "掲載ページ", "照合の結果"])
        w.writerows(results)
    ok = sum(1 for x in results if x[-1] == "OK")
    print(f"照合：{len(results)}行のうち OK {ok}行。結果は {os.path.relpath(out, ROOT)}")
    for x in results:
        if x[-1] != "OK":
            print(f"  {x[1]} {x[2]}行目 {x[3]} {x[4]} {x[5]} 「{x[6]}」 {x[8]} → {x[9]}")


if __name__ == "__main__":
    main()
