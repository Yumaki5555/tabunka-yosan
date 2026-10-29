# 自治体マスタ・年度別人口統計・地図の形を作るプログラム
#
#   python tools/make_master.py
#
# 作るもの
#   data/master/municipalities.csv  … 47都道府県と全市区町村（総務省「全国地方公共団体コード」令和6年1月1日現在）
#   data/master/stat_sources.csv    … 人口・外国人住民数の出典（e-Stat の資料）
#   data/population.csv             … 年度ごとの人口（住民基本台帳）と外国人住民数（在留外国人統計）
#   data/geo/japan.json・pref-*.json … 地図の形（「数でみる在留外国人」と同じもの）
#
# 元にするもの
#   tools/source/soumu_dantai_code_R6-01-01.xlsx（総務省 https://www.soumu.go.jp/denshijiti/code.html ）
#   「数でみる在留外国人」の作業フォルダ（../ビザ）にある、e-Stat から作った表
#     データ/人口/住民基本台帳人口.csv（総務省「住民基本台帳に基づく人口、人口動態及び世帯数調査」）
#     site/data/muni-total.json・prefs.json・meta.json（出入国在留管理庁「在留外国人統計」）
#
# 年度と統計の時点の決まり（画面にも表示します）
#   人口      … その年度の1月1日現在の住民基本台帳人口（例：2026年度 → 2026年1月1日）
#   外国人住民 … その年度の前年12月末の在留外国人数（例：2026年度 → 2025年12月末）

import csv
import json
import os
import shutil
import sys
import unicodedata
from urllib.parse import quote

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ZAIRYU = os.environ.get("ZAIRYU_DIR", os.path.join(ROOT, "..", "ビザ"))
XLSX = os.path.join(ROOT, "tools", "source", "soumu_dantai_code_R6-01-01.xlsx")
SOUMU_URL = "https://www.soumu.go.jp/denshijiti/code.html"
ZAIRYU_SITE = "https://yumaki5555.github.io/zairyu/"
FIRST_FY, LAST_FY = 2016, 2026


def hira(s):
    """半角カナ → 全角ひらがな（検索用）"""
    s = unicodedata.normalize("NFKC", s or "")
    return "".join(chr(ord(c) - 0x60) if "ァ" <= c <= "ヶ" else c for c in s)


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    print(f"  {os.path.relpath(path, ROOT)}  {len(rows)}行")


def main():
    if not os.path.isdir(ZAIRYU):
        sys.exit(f"「数でみる在留外国人」の作業フォルダが見つかりません: {ZAIRYU}（環境変数 ZAIRYU_DIR で指定できます）")

    wb = openpyxl.load_workbook(XLSX, read_only=True)
    sheets = wb.worksheets
    rows = [r for r in sheets[0].iter_rows(values_only=True)][1:]
    seirei = {str(r[0]) for r in list(sheets[1].iter_rows(values_only=True))[1:] if r[0] and r[2] and not str(r[2]).endswith("区")}

    total = json.load(open(os.path.join(ZAIRYU, "site", "data", "muni-total.json"), encoding="utf-8"))
    prefs_stat = json.load(open(os.path.join(ZAIRYU, "site", "data", "prefs.json"), encoding="utf-8"))
    meta = json.load(open(os.path.join(ZAIRYU, "site", "data", "meta.json"), encoding="utf-8"))
    roma_of_pref = {}
    for key, v in total.items():
        if not key.startswith("city:"):
            roma_of_pref[key[:2]] = v[0]
    zairyu_pages = set(os.listdir(os.path.join(ZAIRYU, "site")))

    geo_dir = os.path.join(ZAIRYU, "site_assets", "map")
    geo_keys = {}
    for roma in set(roma_of_pref.values()):
        g = json.load(open(os.path.join(geo_dir, f"pref-{roma}.json"), encoding="utf-8"))
        geo_keys[roma] = set(g["areas"].keys())

    # ---------- 自治体マスタ ----------
    master = []
    for r in rows:
        code, pref, name, pref_kana, kana = [(str(x).strip() if x is not None else "") for x in r[:5]]
        if not code or not code.isdigit():
            continue
        code5 = code[:5]
        if "01695" <= code5 <= "01700":
            # 北方領土の6村（色丹村・泊村・留夜別村・留別村・紗那村・蘂取村）はコードはあるが役場がないため対象外
            continue
        pc = code[:2]
        roma = roma_of_pref.get(pc, "")
        if not name:
            kind, name, kana = "都道府県", pref, pref_kana
            geo_key = pref
            zairyu = f"pref-{roma}.html"
        else:
            if code in seirei:
                kind = "政令指定都市"
            elif pc == "13" and name.endswith("区"):
                kind = "特別区"
            elif name.endswith("市"):
                kind = "市"
            elif name.endswith("町"):
                kind = "町"
            elif name.endswith("村"):
                kind = "村"
            else:
                kind = "その他"
            geo_key = f"city:{name}" if kind == "政令指定都市" else code5
            if geo_key not in geo_keys.get(roma, set()):
                geo_key = ""
            zairyu = f"city-{name}.html" if kind == "政令指定都市" else f"muni-{code5}.html"
        if zairyu not in zairyu_pages:
            zairyu = ""
        master.append([code, code5, pc, pref, name, hira(kana), kind, roma, geo_key, (ZAIRYU_SITE + quote(zairyu)) if zairyu else ""])

    write_csv(os.path.join(ROOT, "data", "master", "municipalities.csv"),
              ["code", "code5", "pref_code", "pref", "name", "kana", "kind", "geo_file", "geo_key", "zairyu_url"], master)
    missing_geo = [m for m in master if m[6] != "都道府県" and not m[8]]
    if missing_geo:
        print("  地図の形が見つからない自治体:", " ".join(m[3] + m[4] for m in missing_geo))

    # ---------- 統計の出典 ----------
    stat_sources = []
    pop_src = {}
    with open(os.path.join(ZAIRYU, "元データ", "人口_出どころ一覧.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            sid = f"POP-{r['年']}"
            pop_src[r["年"]] = sid
            stat_sources.append([sid, f"住民基本台帳に基づく人口、人口動態及び世帯数調査 {r['表題']}（{r['基準日']}現在）",
                                 "総務省", r["URL"], r["公開日"], r["基準日"], "人口（日本人と外国人の合計）"])
    for_src = {}
    with open(os.path.join(ZAIRYU, "元データ", "出どころ一覧.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            p = r["時点"]
            if not p.endswith("-12") or p in for_src or "市区町村別" not in r["表題"]:
                continue
            if int(p[:4]) < FIRST_FY - 1:
                continue
            sid = f"FOR-{p}"
            for_src[p] = sid
            stat_sources.append([sid, f"在留外国人統計 {r['表題']}（{p[:4]}年12月末）", "出入国在留管理庁",
                                 r["URL"], r["公開日"], f"{p[:4]}-12-31", "在留外国人数（中長期在留者＋特別永住者）"])
    write_csv(os.path.join(ROOT, "data", "master", "stat_sources.csv"),
              ["source_id", "title", "publisher", "url", "published_date", "as_of", "note"], stat_sources)

    # ---------- 年度別人口統計 ----------
    pop = {}
    with open(os.path.join(ZAIRYU, "データ", "人口", "住民基本台帳人口.csv"), encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            pop[(r["団体コード"], int(r["基準日"][:4]))] = int(r["人口"])
    MP, P = meta["MP"], meta["P"]
    out = []
    for m in master:
        code, code5, pc, pref, name, kana, kind, roma, geo_key, _ = m
        if kind == "都道府県":
            series, periods = prefs_stat["t"].get(pref), P
        else:
            series, periods = total.get(f"city:{name}" if kind == "政令指定都市" else code5, [None, None])[1], MP
        for fy in range(FIRST_FY, LAST_FY + 1):
            p_val = pop.get((code5, fy))
            fp = f"{fy - 1}-12"
            f_val = series[periods.index(fp)] if series and fp in periods else None
            if p_val is None and f_val is None:
                continue
            out.append([code, fy,
                        p_val if p_val is not None else "", f"{fy}-01-01" if p_val is not None else "",
                        pop_src.get(str(fy), "") if p_val is not None else "",
                        f_val if f_val is not None else "", f"{fy - 1}-12-31" if f_val is not None else "",
                        for_src.get(fp, "") if f_val is not None else ""])
    write_csv(os.path.join(ROOT, "data", "population.csv"),
              ["code", "fy", "population", "population_as_of", "population_source", "foreign", "foreign_as_of", "foreign_source"], out)

    # ---------- 地図の形 ----------
    dst = os.path.join(ROOT, "data", "geo")
    os.makedirs(dst, exist_ok=True)
    for fn in os.listdir(geo_dir):
        if fn == "japan.json" or fn.startswith("pref-"):
            shutil.copyfile(os.path.join(geo_dir, fn), os.path.join(dst, fn))
    print(f"  data/geo/ に地図の形 {len([f for f in os.listdir(dst) if f.endswith('.json')])} ファイル")


if __name__ == "__main__":
    main()
